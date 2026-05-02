"""Tests for the ImportExcelToRevit-compatible Excel exporter.

Pure-function tests — no DB fixture. Covers:
* sanitize_sheet_name truncation + bad-char replacement + empty-name fallback
* collect_columns ordering (id_field first, extras next, row keys after)
* build_excel writes header + rows correctly via openpyxl roundtrip
* add_string_type_annotation appends " : String" everywhere except id_field
* new_-prefixed columns are preserved (the add-in keys on that prefix)
* is_new_parameter_column returns True only for the new_ prefix family
"""

from __future__ import annotations

from openpyxl import load_workbook

from app.modules.ddc_revit_export.service import (
    build_excel,
    collect_columns,
    is_new_parameter_column,
    sanitize_sheet_name,
)


# ── sanitize_sheet_name ────────────────────────────────────────────────────


def test_sanitize_sheet_name_truncates_to_25():
    name = sanitize_sheet_name("a" * 50)
    assert len(name) == 25
    assert name == "a" * 25


def test_sanitize_sheet_name_strips_forbidden_chars():
    assert sanitize_sheet_name("Project / Phase 2: BoQ") == "Project _ Phase 2_ BoQ"


def test_sanitize_sheet_name_handles_empty_after_strip():
    assert sanitize_sheet_name("") == "Model"
    assert sanitize_sheet_name("   ") == "Model"


# ── collect_columns ────────────────────────────────────────────────────────


def test_collect_columns_id_first_then_extras_then_row_keys():
    rows = [
        {"ElementId": 1, "Mark": "W1", "Comments": "x"},
        {"ElementId": 2, "Mark": "W2", "FireRating": "F90"},
    ]
    cols = collect_columns(rows, id_field="ElementId", extra_columns=["new_AssetTag"])
    assert cols == ["ElementId", "new_AssetTag", "Mark", "Comments", "FireRating"]


def test_collect_columns_no_duplicates_when_extra_already_in_rows():
    rows = [{"ElementId": 1, "Mark": "W1"}]
    cols = collect_columns(rows, id_field="ElementId", extra_columns=["Mark"])
    assert cols == ["ElementId", "Mark"]


# ── build_excel roundtrip ──────────────────────────────────────────────────


def test_build_excel_writes_header_and_rows():
    rows = [
        {"ElementId": 1001, "Mark": "W1", "Comments": "exterior"},
        {"ElementId": 1002, "Mark": "W2", "Comments": ""},
    ]
    buf = build_excel("Project Alpha", rows)
    wb = load_workbook(buf)
    ws = wb.active
    assert ws.title == "Project Alpha"
    header = [c.value for c in ws[1]]
    assert header == ["ElementId", "Mark", "Comments"]
    assert [c.value for c in ws[2]] == [1001, "W1", "exterior"]
    # openpyxl normalises empty-string cells back to None on read
    assert [c.value for c in ws[3]] == [1002, "W2", None]


def test_build_excel_truncates_long_model_name_in_sheet():
    buf = build_excel("M" * 50, [{"ElementId": 1}])
    wb = load_workbook(buf)
    assert wb.active.title == "M" * 25


def test_build_excel_string_type_annotation_only_on_non_id_columns():
    rows = [{"ElementId": 1, "Mark": "W1", "Comments": "x"}]
    buf = build_excel("M", rows, add_string_type_annotation=True)
    wb = load_workbook(buf)
    header = [c.value for c in wb.active[1]]
    assert header == ["ElementId", "Mark : String", "Comments : String"]


def test_build_excel_preserves_new_prefix_columns():
    rows = [{"ElementId": 1, "Mark": "W1", "new_AssetTag": "AT-001"}]
    buf = build_excel("M", rows)
    wb = load_workbook(buf)
    header = [c.value for c in wb.active[1]]
    assert "new_AssetTag" in header
    assert wb.active.cell(row=2, column=header.index("new_AssetTag") + 1).value == "AT-001"


def test_build_excel_handles_none_as_empty_cell():
    rows = [{"ElementId": 1, "Mark": None}]
    buf = build_excel("M", rows)
    wb = load_workbook(buf)
    assert wb.active.cell(row=2, column=2).value is None


# ── is_new_parameter_column ────────────────────────────────────────────────


def test_is_new_parameter_column_detects_prefix():
    assert is_new_parameter_column("new_AssetTag") is True
    assert is_new_parameter_column("new_AssetTag : String") is True


def test_is_new_parameter_column_rejects_other_columns():
    assert is_new_parameter_column("Mark") is False
    assert is_new_parameter_column("Comments") is False
    assert is_new_parameter_column("ElementId") is False
    # case-sensitive: NEW_ doesn't trigger (matches the add-in's strict prefix check)
    assert is_new_parameter_column("NEW_AssetTag") is False
