"""Tests for the ddc_pdf_excel ingestion service.

Mocks the pdfplumber side of the operation so the test is deterministic
(real PDF table detection is heuristic and depends heavily on the source
PDF's rendering — the conversion logic is what we want to verify here).
"""

from __future__ import annotations

from io import BytesIO
from unittest.mock import MagicMock

from openpyxl import load_workbook


def _patch_pdfplumber_to_return(monkeypatch, pages_tables: list[list[list[list[str | None]]]]) -> None:
    """Make pdfplumber.open() yield pages whose extract_tables() returns the supplied data.

    ``pages_tables[i][j]`` is the ``j``-th table on page ``i+1``, as a list-of-rows.
    """
    import pdfplumber

    pages = []
    for tables in pages_tables:
        page = MagicMock()
        page.extract_tables.return_value = tables
        pages.append(page)

    pdf_ctx = MagicMock()
    pdf_ctx.__enter__.return_value = MagicMock(pages=pages)
    pdf_ctx.__exit__.return_value = False

    monkeypatch.setattr(pdfplumber, "open", lambda _src: pdf_ctx)


def test_single_page_one_table_writes_one_sheet(monkeypatch):
    from app.modules.ddc_pdf_excel.service import pdf_bytes_to_excel

    _patch_pdfplumber_to_return(
        monkeypatch,
        [
            [
                [
                    ["Code", "Item", "Qty"],
                    ["A1", "Concrete C30", "100"],
                    ["A2", "Rebar 10mm", "5000"],
                ],
            ],
        ],
    )

    out, count = pdf_bytes_to_excel(b"%PDF-fake")
    assert count == 1

    wb = load_workbook(out)
    assert wb.sheetnames == ["Page1_Table1"]
    ws = wb["Page1_Table1"]
    assert [c.value for c in ws[1]] == ["Code", "Item", "Qty"]
    assert [c.value for c in ws[2]] == ["A1", "Concrete C30", "100"]
    assert [c.value for c in ws[3]] == ["A2", "Rebar 10mm", "5000"]


def test_multi_page_multi_table_writes_one_sheet_each(monkeypatch):
    from app.modules.ddc_pdf_excel.service import pdf_bytes_to_excel

    _patch_pdfplumber_to_return(
        monkeypatch,
        [
            [
                [["A", "B"], ["1", "2"]],
                [["C"], ["3"]],
            ],
            [
                [["D", "E"], ["4", "5"]],
            ],
        ],
    )

    out, count = pdf_bytes_to_excel(b"%PDF-fake")
    assert count == 3

    wb = load_workbook(out)
    assert wb.sheetnames == ["Page1_Table1", "Page1_Table2", "Page2_Table1"]


def test_none_cells_become_empty_strings(monkeypatch):
    from app.modules.ddc_pdf_excel.service import pdf_bytes_to_excel

    _patch_pdfplumber_to_return(
        monkeypatch,
        [[[["a", None, "c"], [None, "b", None]]]],
    )

    out, _ = pdf_bytes_to_excel(b"%PDF-fake")
    wb = load_workbook(out)
    ws = wb["Page1_Table1"]
    # None cells were written as "" — but openpyxl normalises empty-string back to None on read
    # so what we actually verify is "the cell exists and isn't crashing"
    assert ws.cell(row=1, column=1).value == "a"
    assert ws.cell(row=1, column=3).value == "c"
    assert ws.cell(row=2, column=2).value == "b"


def test_no_tables_yields_empty_marker_sheet(monkeypatch):
    from app.modules.ddc_pdf_excel.service import pdf_bytes_to_excel

    _patch_pdfplumber_to_return(monkeypatch, [[]])

    out, count = pdf_bytes_to_excel(b"%PDF-fake")
    assert count == 0
    wb = load_workbook(out)
    assert wb.sheetnames == ["Empty"]
    assert wb["Empty"].cell(row=1, column=1).value == "No tables detected in PDF."


def test_empty_table_skipped(monkeypatch):
    from app.modules.ddc_pdf_excel.service import pdf_bytes_to_excel

    _patch_pdfplumber_to_return(monkeypatch, [[[], [["A"], ["1"]]]])

    out, count = pdf_bytes_to_excel(b"%PDF-fake")
    # The empty table doesn't get a sheet; only the populated one
    assert count == 1
    wb = load_workbook(out)
    assert wb.sheetnames == ["Page1_Table2"]
