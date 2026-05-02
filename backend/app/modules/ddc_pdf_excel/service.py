"""Extract tables from a PDF and write them to a multi-sheet .xlsx.

pdfplumber returns tables as ``list[list[Cell|None]]`` where each inner
list is a row. We:

* iterate pages
* extract each page's tables (default ``extract_tables`` settings)
* write one Excel sheet per table — sheet names ``Page<N>_Table<M>``,
  truncated to Excel's 31-char limit
* write the table's first row as-is (assumed header) and subsequent rows
  as data — no further normalisation, since DDC's notebook didn't do
  any either

The function returns a ``BytesIO`` ready to stream. If pdfplumber finds
no tables on any page, the workbook contains a single sheet named
``Empty`` with a one-cell explanation.
"""

from __future__ import annotations

import re
from io import BytesIO

_BAD_SHEET_CHARS = re.compile(r"[\[\]:\*\?/\\]")
_MAX_SHEET_NAME_LEN = 31  # Excel hard limit on sheet names


def _sanitize_sheet(name: str) -> str:
    cleaned = _BAD_SHEET_CHARS.sub("_", name).strip()
    if not cleaned:
        cleaned = "Sheet"
    return cleaned[:_MAX_SHEET_NAME_LEN]


def pdf_bytes_to_excel(pdf_bytes: bytes) -> tuple[BytesIO, int]:
    """Convert PDF bytes to a multi-sheet .xlsx (one sheet per table).

    Returns:
        ``(BytesIO, table_count)``. Caller can use the table count to
        report a 422 if no tables were found.
    """
    try:
        import pdfplumber
    except ImportError as e:  # pragma: no cover
        raise RuntimeError(
            "pdfplumber required for PDF ingestion (it's already a NEXUS base dep)"
        ) from e
    try:
        from openpyxl import Workbook
    except ImportError as e:  # pragma: no cover
        raise RuntimeError("openpyxl required for Excel writing") from e

    wb = Workbook()
    # Remove the default sheet so we can append clean
    default_sheet = wb.active
    wb.remove(default_sheet)

    table_count = 0
    used_names: set[str] = set()

    with pdfplumber.open(BytesIO(pdf_bytes)) as pdf:
        for page_idx, page in enumerate(pdf.pages, start=1):
            tables = page.extract_tables() or []
            for tbl_idx, table in enumerate(tables, start=1):
                if not table:
                    continue
                table_count += 1
                base_name = _sanitize_sheet(f"Page{page_idx}_Table{tbl_idx}")
                # de-dupe in the unlikely case sanitization collides
                name = base_name
                suffix = 2
                while name in used_names:
                    name = f"{base_name[: _MAX_SHEET_NAME_LEN - 3]}_{suffix}"
                    suffix += 1
                used_names.add(name)
                ws = wb.create_sheet(title=name)
                for row in table:
                    ws.append([(c if c is not None else "") for c in row])

    if table_count == 0:
        ws = wb.create_sheet(title="Empty")
        ws.append(["No tables detected in PDF."])

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out, table_count
